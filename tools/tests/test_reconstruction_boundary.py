import contextlib
import io
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from deckit_dev.cli import (
    _audit_pptx_container,
    _find_pptx_deliverables,
    _scan_image_first_artifacts,
    new_run,
)


class ReconstructionBoundaryTests(unittest.TestCase):
    def test_defaults_and_premium_do_not_request_reconstruction(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.md"
            source.write_text("Test deck source", encoding="utf-8")
            for budget in (None, "premium"):
                name = budget or "default"
                with contextlib.redirect_stdout(io.StringIO()):
                    new_run(str(source), name, root / "outputs", False,
                            None, budget, "pptx", 1, None)
                manifest = json.loads((root / "outputs" / name / "run.json").read_text())
                self.assertEqual(manifest["production_mode"], "image-first")
                self.assertEqual(manifest["native_conversion"], {
                    "route": "none", "status": "not-requested", "slide_ids": [],
                })

    def make_deliveries(self, root):
        generated = root / "assets" / "generated-slides"
        generated.mkdir(parents=True)
        image = generated / "00_cover.png"
        Image.new("RGB", (1600, 900), "white").save(image)
        dist = root / "dist"
        dist.mkdir()
        Image.new("RGB", (900, 506), "white").save(dist / "preview.png")

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(image), 0, 0,
                                 width=prs.slide_width, height=prs.slide_height)
        baseline = dist / "bitmap.pptx"
        prs.save(baseline)

        native_dir = dist / "native" / "conversion-01"
        native_dir.mkdir(parents=True)
        native = Presentation()
        slide = native.slides.add_slide(native.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Editable title"
        derivative = native_dir / "candidate.pptx"
        native.save(derivative)
        return baseline, derivative

    def test_nested_derivative_does_not_change_baseline_audit(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            baseline, derivative = self.make_deliveries(root)
            self.assertEqual(_find_pptx_deliverables(root), [baseline])
            findings = _scan_image_first_artifacts(root, ["00_cover"])
            self.assertFalse([f for f in findings if f[0] == "error"], findings)
            # A separate path is not native acceptance: strict audit still rejects it.
            findings, _ = _audit_pptx_container(derivative)
            self.assertIn("PPTX-NATIVE-CONTENT", [f[1] for f in findings])

    def test_native_file_in_baseline_directory_still_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            _, derivative = self.make_deliveries(root)
            shutil.copy2(derivative, root / "dist" / "unexpected-native.pptx")
            findings = _scan_image_first_artifacts(root, ["00_cover"])
            self.assertIn("PPTX-NATIVE-CONTENT", [f[1] for f in findings])


if __name__ == "__main__":
    unittest.main()
