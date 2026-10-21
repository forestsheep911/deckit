from __future__ import annotations

import contextlib
import io
import json
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation

from deckit_dev.cli import _audit_pptx_container, _package_images, _package_preview
from deckit_dev.flow_artifacts import scan_flow_artifacts


REPO = Path(__file__).resolve().parents[2]
SKILL = REPO / "plugins" / "deckit" / "skills" / "complex-flow"
SAMPLE = SKILL / "assets" / "approval.workflow.json"
SCRIPT = SKILL / "scripts" / "render-flow.mjs"
FLOW_BODY = "- **Visual method**: complex-flow\n- **Routing reason**: Three roles and a rejection/retry loop.\n- **Speaker notes**: Explain approval and resubmission.\n"


@unittest.skipUnless(shutil.which("node"), "Node.js required")
class ComplexFlowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.run_dir = Path(cls.tmp.name)
        cls.work = cls.run_dir / "work" / "diagrams" / "01_approval_flow"
        cls.work.mkdir(parents=True)
        cls.source = cls.work / "workflow.json"
        shutil.copyfile(SAMPLE, cls.source)
        cls.command = ["node", str(SCRIPT), "--run-dir", str(cls.run_dir), "--slide-id", "01_approval_flow",
                       "--input", str(cls.source), "--reason", "Three roles and a rejection/retry loop."]
        result = subprocess.run(cls.command, capture_output=True, text=True, encoding="utf-8", timeout=60)
        if result.returncode:
            cls.tmp.cleanup()
            raise AssertionError(result.stderr)
        cls.receipt = json.loads(result.stdout)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_real_capture_keeps_topology_and_exact_viewport(self):
        receipt = self.receipt
        source = json.loads(self.source.read_text(encoding="utf-8"))
        self.assertEqual({(e["from"], e["to"]) for e in receipt["layout"]["edges"]},
                         {(e["from"], e["to"]) for e in source["edges"]})
        self.assertTrue(receipt["complexity"]["cycle"])
        self.assertEqual((receipt["browser_check"]["width"], receipt["browser_check"]["height"]), (1600, 900))
        self.assertGreaterEqual(receipt["browser_check"]["minNodePx"], 18)
        with Image.open(self.run_dir / receipt["png"]) as image:
            image.load()
            self.assertEqual(image.size, (3200, 1800))
        self.assertEqual(scan_flow_artifacts(self.run_dir, [("01_approval_flow", FLOW_BODY)]), [])

    def test_mixed_image_container_keeps_order_and_notes(self):
        images = self.run_dir / "assets" / "generated-slides"
        # Fixtures stand for imagegen pages; this test does not claim an imagegen call.
        for sid, color in [("00_cover", "navy"), ("02_closing", "white")]:
            Image.new("RGB", (1600, 900), color).save(images / f"{sid}.png")
        (self.run_dir / "work" / "storyboard.md").write_text(
            "## 00_cover\n- **Speaker notes**: Cover fixture.\n\n"
            "## 01_approval_flow\n" + FLOW_BODY + "\n## 02_closing\n- **Speaker notes**: Closing fixture.\n", encoding="utf-8")
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(_package_images(self.run_dir, None), 0)
            _package_preview(self.run_dir, None)
        pptx = self.run_dir / "dist" / f"{self.run_dir.name}.pptx"
        findings, _ = _audit_pptx_container(pptx)
        self.assertFalse([f for f in findings if f[0] == "error"], findings)
        slides = Presentation(pptx).slides
        self.assertEqual(len(slides), 3)
        self.assertEqual(slides[1].shapes[0].image.blob, (images / "01_approval_flow.png").read_bytes())
        self.assertIn("approval and resubmission", slides[1].notes_slide.notes_text_frame.text)
        self.assertTrue((self.run_dir / "dist" / "preview.png").is_file())

    def test_changed_source_and_unselected_route_are_rejected(self):
        original = self.source.read_bytes()
        try:
            self.source.write_bytes(original + b"\n")
            findings = scan_flow_artifacts(self.run_dir, [("01_approval_flow", FLOW_BODY)])
            self.assertTrue(any(code == "FLOW-ARTIFACT" for _, code, _ in findings))
        finally:
            self.source.write_bytes(original)
        self.assertTrue(any(code == "FLOW-UNSELECTED" for _, code, _ in scan_flow_artifacts(self.run_dir, [])))

    def test_existing_png_is_preserved(self):
        png = self.run_dir / self.receipt["png"]
        before = png.read_bytes()
        result = subprocess.run(self.command, capture_output=True, text=True, encoding="utf-8", timeout=20)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("already exists", result.stderr)
        self.assertEqual(png.read_bytes(), before)

    def test_simple_linear_flow_does_not_start_browser(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = json.loads(SAMPLE.read_text(encoding="utf-8"))
            source["lanes"] = source["lanes"][:1]
            source["nodes"] = source["nodes"][:2]
            for node in source["nodes"]:
                node["lane"] = "requester"
            source["edges"] = source["edges"][:1]
            candidate = root / "simple.json"
            candidate.write_text(json.dumps(source), encoding="utf-8")
            result = subprocess.run(["node", str(SCRIPT), "--run-dir", tmp, "--slide-id", "00_simple",
                                     "--input", str(candidate), "--reason", "Simple", "--browser", "missing-browser"],
                                    capture_output=True, text=True, encoding="utf-8", timeout=20)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("simple linear flow", result.stderr)
            self.assertFalse((root / "assets").exists())

    def test_broken_topology_and_output_escape_fail(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            candidate = root / "broken.json"
            source = json.loads(SAMPLE.read_text(encoding="utf-8"))
            source["edges"][0]["to"] = "missing_node"
            candidate.write_text(json.dumps(source), encoding="utf-8")
            args = ["node", str(SCRIPT), "--run-dir", tmp, "--slide-id", "01_broken", "--input", str(candidate), "--reason", "Retry loop"]
            result = subprocess.run(args, capture_output=True, text=True, encoding="utf-8", timeout=20)
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse((root / "assets").exists())
            args[args.index(str(candidate))] = str(SAMPLE)
            result = subprocess.run(args, capture_output=True, text=True, encoding="utf-8", timeout=20)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("child of the run", result.stderr)

    def test_missing_browser_is_explicit_and_publishes_no_png(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            candidate = root / "workflow.json"
            shutil.copyfile(SAMPLE, candidate)
            result = subprocess.run(["node", str(SCRIPT), "--run-dir", tmp, "--slide-id", "01_flow",
                                     "--input", str(candidate), "--reason", "Retry loop", "--browser", str(root / "missing-browser")],
                                    capture_output=True, text=True, encoding="utf-8", timeout=20)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("Requested browser executable", result.stderr)
            self.assertFalse((root / "assets").exists())

    def test_actual_title_overflow_never_publishes_a_png(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = json.loads(SAMPLE.read_text(encoding="utf-8"))
            source["meta"]["title"] = "审批说明必须清晰可读" * 30
            candidate = root / "workflow.json"
            candidate.write_text(json.dumps(source), encoding="utf-8")
            result = subprocess.run(["node", str(SCRIPT), "--run-dir", tmp, "--slide-id", "01_flow",
                                     "--input", str(candidate), "--reason", "Retry loop"],
                                    capture_output=True, text=True, encoding="utf-8", timeout=40)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("Slide readability failed", result.stderr)
            self.assertFalse((root / "assets" / "generated-slides" / "01_flow.png").exists())


if __name__ == "__main__":
    unittest.main()
